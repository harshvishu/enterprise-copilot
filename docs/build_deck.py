"""Generate the Enterprise Copilot workshop deck (architecture + Confluence agent + speaker notes).

Run: python docs/build_deck.py  (requires python-pptx)
Output: docs/enterprise-copilot-briefing.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Palette (dark workshop theme)
BG = RGBColor(0x11, 0x16, 0x27)
CARD = RGBColor(0x1B, 0x24, 0x36)
ACCENT = RGBColor(0xF2, 0xB5, 0x44)
TEXT = RGBColor(0xE7, 0xEC, 0xF5)
MUTED = RGBColor(0x9A, 0xA6, 0xBF)
GREEN = RGBColor(0x4C, 0xC9, 0x8F)

EMU_W, EMU_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


def bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def box(slide, l, t, w, h, fill=None, line=None):
    shp = slide.shapes.add_shape(1, l, t, w, h)  # rectangle
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill if fill else CARD
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def text(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (txt, size, color, bold) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = "Segoe UI"
    return tb


def bullets(slide, l, t, w, h, items, size=16, color=TEXT, space=8):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (txt, lvl, c, b) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = ("• " if lvl == 0 else "– ") + txt
        r.font.size = Pt(size)
        r.font.color.rgb = c or color
        r.font.bold = b
        r.font.name = "Segoe UI"
    return tb


def notes(slide, note):
    slide.notes_slide.notes_text_frame.text = note


def header(slide, kicker, title):
    box(slide, 0, 0, EMU_W, Inches(0.12), fill=ACCENT)
    text(slide, Inches(0.6), Inches(0.35), Inches(12), Inches(0.4),
         [(kicker, 14, ACCENT, True)])
    text(slide, Inches(0.6), Inches(0.7), Inches(12.1), Inches(0.9),
         [(title, 30, TEXT, True)])


# ---------------------------------------------------------------- Slide 1: title
s = prs.slides.add_slide(BLANK)
bg(s)
box(s, 0, Inches(2.6), EMU_W, Inches(0.05), fill=ACCENT)
text(s, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.0),
     [("Enterprise Copilot", 48, TEXT, True)])
text(s, Inches(0.8), Inches(2.75), Inches(11.7), Inches(0.6),
     [("Architecture · Confluence Agent · Speaker Notes", 22, MUTED, False)])
text(s, Inches(0.8), Inches(4.4), Inches(11.7), Inches(0.8),
     [("\u201cAI proposes. Humans dispose.\u201d", 26, ACCENT, True)])
text(s, Inches(0.8), Inches(5.2), Inches(11.7), Inches(0.6),
     [("AI accelerates delivery; humans own accountability.", 16, MUTED, False)])
text(s, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.4),
     [("Flo 2026 — Spring AI Workshop", 14, MUTED, False)])
notes(s, "Meet four AI teammates and the human who holds the deploy button. This deck covers how the "
         "system is built, how to add the Confluence Agent exercise, and how to present the demo.")

# ---------------------------------------------------------------- Slide 2: architecture
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "1 · ARCHITECTURE", "Modular monolith + React dashboard")
bullets(s, Inches(0.6), Inches(1.8), Inches(6.1), Inches(5),
        [("Spring Boot modular monolith — one process, clear boundaries", 0, TEXT, True),
         ("api/ controllers expose REST + SSE", 1, MUTED, False),
         ("orchestration/ PipelineOrchestrator drives the flow", 1, MUTED, False),
         ("agents/ Rhea, Nova, Sentinel (LLM)", 1, MUTED, False),
         ("Atlas — deterministic Java gates", 1, MUTED, False),
         ("tools/ Compliance, Architecture, GitHistory, ApiSpec, Confluence", 1, MUTED, False),
         ("React + Vite + Tailwind dashboard", 0, TEXT, True),
         ("REST + SSE; UI auto-detects optional stages from events", 1, MUTED, False),
         ("Persistence", 0, TEXT, True),
         ("pipelines snapshot (JSON) + append-only redacted audit_events", 1, MUTED, False),
         ("Flyway owns schema; Hibernate validate-only", 1, MUTED, False)], size=15)
box(s, Inches(7.0), Inches(1.9), Inches(5.7), Inches(4.7), fill=CARD)
text(s, Inches(7.2), Inches(2.05), Inches(5.3), Inches(0.4),
     [("Flow", 16, ACCENT, True)])
bullets(s, Inches(7.2), Inches(2.55), Inches(5.3), Inches(4),
        [("UI  →  API  →  Orchestrator", 0, TEXT, True),
         ("Orchestrator → Agents → Tools", 0, TEXT, True),
         ("Agents → AgentAiClient → OpenAI / Ollama", 0, TEXT, True),
         ("Orchestrator → Atlas (deterministic gates)", 0, TEXT, True),
         ("Orchestrator → AuditService → Redactor", 0, TEXT, True),
         ("Orchestrator → Persistence → H2 / PostgreSQL", 0, TEXT, True)], size=15, space=12)
notes(s, "One process, no premature microservices. The orchestrator coordinates three model-backed "
         "agents and one deterministic release manager. Everything is auditable and redacted.")

# ---------------------------------------------------------------- Slide 3: four agents
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "1 · ARCHITECTURE", "Four agents, like new teammates")
cards = [
    ("\U0001F50D Rhea", "Requirements Analyst", "LLM", "Refuses to guess. Raises clarification questions that PAUSE the pipeline."),
    ("\U0001F4BB Nova", "Senior Java Engineer", "LLM", "Proposes code as a PR diff + tests. Never touches the filesystem."),
    ("\U0001F6E1 Sentinel", "Security & Compliance", "LLM", "Severity-tagged findings. Any CRITICAL blocks deploy. The hero moment."),
    ("\U0001F680 Atlas", "Release Manager", "Deterministic Java", "Rule-based gates. Allowed only when every gate passes AND a human approves."),
]
x = Inches(0.6)
w = Inches(2.95)
gap = Inches(0.15)
for i, (name, role, kind, desc) in enumerate(cards):
    lx = x + i * (w + gap)
    kcolor = GREEN if "Deterministic" in kind else ACCENT
    box(s, lx, Inches(1.9), w, Inches(4.6), fill=CARD)
    text(s, lx + Inches(0.15), Inches(2.05), w - Inches(0.3), Inches(0.6), [(name, 20, TEXT, True)])
    text(s, lx + Inches(0.15), Inches(2.7), w - Inches(0.3), Inches(0.5), [(role, 13, MUTED, False)])
    box(s, lx + Inches(0.15), Inches(3.2), Inches(1.6), Inches(0.35), fill=BG)
    text(s, lx + Inches(0.15), Inches(3.2), Inches(1.6), Inches(0.35), [(kind, 11, kcolor, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, lx + Inches(0.15), Inches(3.75), w - Inches(0.3), Inches(2.6), [(desc, 14, TEXT, False)])
notes(s, "Name them like people. Rhea clarifies, Nova proposes, Sentinel catches the flaw, Atlas holds "
         "the gate. Atlas is deterministic Java in every mode — the release decision is never a probability.")

# ---------------------------------------------------------------- Slide 4: pipeline / modes
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "1 · ARCHITECTURE", "Pipeline & execution modes")
box(s, Inches(0.6), Inches(1.9), Inches(12.1), Inches(0.9), fill=CARD)
text(s, Inches(0.8), Inches(1.9), Inches(11.7), Inches(0.9),
     [("ISSUE → REQUIREMENTS → CODE → REVIEW → DEPLOY  (human approval before DEPLOYED)", 18, ACCENT, True)],
     anchor=MSO_ANCHOR.MIDDLE)
bullets(s, Inches(0.6), Inches(3.1), Inches(6.0), Inches(3.5),
        [("LIVE — OpenAI default", 0, TEXT, True),
         ("Rhea, Nova, Sentinel make real model calls; results vary", 1, MUTED, False),
         ("Ollama is the explicit local LIVE alternative", 1, MUTED, False),
         ("DEMO — deterministic preview/fallback", 0, TEXT, True),
         ("Scripted outcomes for 100% reproducible reveals", 1, MUTED, False),
         ("Credential-free self-checks", 1, MUTED, False)], size=16)
bullets(s, Inches(6.9), Inches(3.1), Inches(5.8), Inches(3.5),
        [("Atlas is deterministic Java in EVERY mode", 0, GREEN, True),
         ("Gates: artifacts, clarification, review APPROVE,", 1, MUTED, False),
         ("no CRITICAL findings, valid test signal, human approval", 1, MUTED, False),
         ("Never promise a deterministic LIVE security finding", 0, ACCENT, True),
         ("Tests & deployment are simulated — nothing ships for real", 1, MUTED, False)], size=16)
notes(s, "The reveal must be reproducible, so the hero moment is shown in DEMO. LIVE uses the real "
         "provider. Atlas never changes: a human always owns the last click.")

# ---------------------------------------------------------------- Slide 5: confluence intro
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "2 · EXERCISE", "Add the Confluence Agent")
text(s, Inches(0.6), Inches(1.75), Inches(12.1), Inches(0.8),
     [("Goal: give Rhea internal business context so she resolves an ambiguity without asking a human.", 18, TEXT, True)])
bullets(s, Inches(0.6), Inches(2.7), Inches(12), Inches(4),
        [("Deterministic — NO LLM. Not every agent needs a model.", 0, ACCENT, True),
         ("Reads one supplied local decision document and passes the text to Rhea transiently", 1, MUTED, False),
         ("No new AI kind, result model, prompt, state, persistent field or feature flag", 1, MUTED, False),
         ("Already provided (do NOT build):", 0, TEXT, True),
         ("ConfluenceTool — reads resources/demo-data/confluence.md (no search)", 1, MUTED, False),
         ("RequirementsAgent.analyze(ctx, businessContext) — Rhea's string overload", 1, MUTED, False),
         ("withConfluenceActivity(...) wrapper emits the activity events", 1, MUTED, False),
         ("DEMO plumbing + React UI that auto-detects the optional Confluence stage", 1, MUTED, False),
         ("Participant work: ~15–20 lines + supplied wiring (5 min expected, 8 max)", 0, GREEN, True)], size=15)
notes(s, "Teaching point: not every agent needs an LLM. Rhea already reasons; she just needed the "
         "missing business context. The tool, overload and plumbing are supplied.")

# ---------------------------------------------------------------- Slide 6: code step
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "2 · EXERCISE", "Step 1 — Create the agent")
text(s, Inches(0.6), Inches(1.7), Inches(12), Inches(0.4),
     [("backend/.../agents/confluence/ConfluenceAgent.java", 14, MUTED, False)])
code = (
    "@Component\n"
    "public class ConfluenceAgent {\n"
    "    private final ConfluenceTool confluenceTool;\n\n"
    "    public ConfluenceAgent(ConfluenceTool confluenceTool) {\n"
    "        this.confluenceTool = confluenceTool;\n"
    "    }\n\n"
    "    public String gatherContext(Ticket ticket) {\n"
    "        return confluenceTool.lookup(ticket.description());\n"
    "    }\n"
    "}"
)
box(s, Inches(0.6), Inches(2.2), Inches(12.1), Inches(4.2), fill=RGBColor(0x0C, 0x10, 0x1C))
cb = s.shapes.add_textbox(Inches(0.85), Inches(2.4), Inches(11.6), Inches(3.9))
cf = cb.text_frame
cf.word_wrap = True
for i, line in enumerate(code.split("\n")):
    p = cf.paragraphs[0] if i == 0 else cf.add_paragraph()
    r = p.add_run()
    r.text = line if line else " "
    r.font.size = Pt(15)
    r.font.name = "Consolas"
    r.font.color.rgb = GREEN if "public String" in line or "@Component" in line else TEXT
notes(s, "About 15-20 lines including imports. Constructor injection of the supplied ConfluenceTool; "
         "gatherContext just calls lookup with the ticket description.")

# ---------------------------------------------------------------- Slide 7: wiring step
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "2 · EXERCISE", "Step 2 — Wire it into the orchestrator")
bullets(s, Inches(0.6), Inches(1.9), Inches(6.0), Inches(4),
        [("Fill the five CONFLUENCE_EXERCISE markers:", 0, ACCENT, True),
         ("Import the agent", 1, MUTED, False),
         ("Field: private final ConfluenceAgent confluenceAgent;", 1, MUTED, False),
         ("Constructor parameter: ConfluenceAgent confluenceAgent,", 1, MUTED, False),
         ("Assignment: this.confluenceAgent = confluenceAgent;", 1, MUTED, False),
         ("Replace the initial Rhea call (right)", 1, MUTED, False)], size=16)
box(s, Inches(6.8), Inches(1.9), Inches(5.9), Inches(2.6), fill=RGBColor(0x0C, 0x10, 0x1C))
cb = s.shapes.add_textbox(Inches(7.0), Inches(2.05), Inches(5.6), Inches(2.3))
cf = cb.text_frame
cf.word_wrap = True
for i, line in enumerate([
    "String businessContext =",
    "  withConfluenceActivity(ctx,",
    "    () -> confluenceAgent",
    "          .gatherContext(ctx.ticket()));",
    "RequirementAnalysis analysis =",
    "  requirementsAgent",
    "    .analyze(ctx, businessContext);"]):
    p = cf.paragraphs[0] if i == 0 else cf.add_paragraph()
    r = p.add_run()
    r.text = line
    r.font.size = Pt(14)
    r.font.name = "Consolas"
    r.font.color.rgb = TEXT
notes(s, "Markers keep the exercise scoped. The supplied withConfluenceActivity wrapper records the "
         "activity events that make the Confluence stage appear in the UI.")

# ---------------------------------------------------------------- Slide 8: before/after
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "2 · EXERCISE", "Step 3 — Before / After (UB-4823)")
box(s, Inches(0.6), Inches(1.9), Inches(5.9), Inches(3.0), fill=CARD)
text(s, Inches(0.8), Inches(2.05), Inches(5.5), Inches(0.5), [("BEFORE (starter)", 18, ACCENT, True)])
bullets(s, Inches(0.8), Inches(2.6), Inches(5.5), Inches(2.1),
        [("Rhea pauses and asks:", 0, TEXT, True),
         ("\u201cWhich notification channel should we use?\u201d", 1, MUTED, False),
         ("Flow: Rhea → Human clarification", 0, TEXT, True)], size=16)
box(s, Inches(6.8), Inches(1.9), Inches(5.9), Inches(3.0), fill=CARD)
text(s, Inches(7.0), Inches(2.05), Inches(5.5), Inches(0.5), [("AFTER (with agent)", 18, GREEN, True)])
bullets(s, Inches(7.0), Inches(2.6), Inches(5.5), Inches(2.1),
        [("Confluence supplies the approved SMS decision", 0, TEXT, True),
         ("No human input needed", 1, MUTED, False),
         ("Confluence → Rhea → Nova → Sentinel → Atlas → Approval", 0, TEXT, True)], size=15)
box(s, Inches(0.6), Inches(5.1), Inches(12.1), Inches(1.2), fill=RGBColor(0x2A, 0x1B, 0x1B))
text(s, Inches(0.8), Inches(5.25), Inches(11.7), Inches(1.0),
     [("Negative control — UB-4825 still BLOCKS: the policy supplies no fraud contract, so Sentinel "
       "blocks the unsupported proposal even with Confluence context.", 15, RGBColor(0xF0, 0x9A, 0x9A), False)],
     anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Reset for the starter, apply for the reference. ./workshop/scripts/confluence-exercise.sh "
         "reset | apply. Both idempotent; apply --force only replaces exercise-owned files.")

# ---------------------------------------------------------------- Slide 9: speaker notes overview
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "3 · SPEAKER NOTES", "Opening hook (10s)")
box(s, Inches(0.6), Inches(1.9), Inches(12.1), Inches(2.6), fill=CARD)
text(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(2.3),
     [("\u201cI want you to meet my four newest teammates. They don't drink coffee, they never argue "
       "in standup, and one of them is about to stop a bug from reaching production. This isn't a "
       "chatbot — it's an engineering team made of AI, with a human holding the deploy button.\u201d",
       20, TEXT, True)], anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(0.6), Inches(4.9), Inches(12), Inches(1),
     [("On screen: the strip ISSUE → REQUIREMENTS → CODE → REVIEW → DEPLOY.", 16, MUTED, False),
      ("\u201cWatch it light up left to right.\u201d  →  click Run Pipeline.", 16, ACCENT, True)])
notes(s, "Set the frame: an engineering team made of AI, with a human owning accountability.")

# ---------------------------------------------------------------- Slide 10: agent speaker lines
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "3 · SPEAKER NOTES", "Walk the four agents")
lines = [
    ("\U0001F50D Rhea", "\u201cGive a junior a vague ticket and they start coding. Rhea asks the awkward questions first — which channel? Has the customer consented? POPIA.\u201d", "Ask: who's shipped the wrong thing because a ticket was vague?"),
    ("\U0001F4BB Nova", "\u201cIt's a proposal — a diff, like a pull request. Nothing is applied. No AI reached into my repo.\u201d", "Open Pull Requests; test signals are proposals, not executed tests."),
    ("\U0001F6E1 Sentinel", "\u201cThis example logs the customer's account number. A rushed reviewer misses it at 5pm Friday. Sentinel doesn't.\u201d", "Pause. Show of hands — who's sure that bug never shipped?"),
    ("\U0001F680 Atlas", "\u201cAtlas is NOT an LLM. The AI can recommend all day. It cannot click this button.\u201d", "Audience vote: should the AI deploy anyway? No — we own it."),
]
y = Inches(1.85)
for name, say, move in lines:
    box(s, Inches(0.6), y, Inches(12.1), Inches(1.15), fill=CARD)
    text(s, Inches(0.8), y + Inches(0.1), Inches(2.0), Inches(0.95), [(name, 18, ACCENT, True)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(2.8), y + Inches(0.08), Inches(9.7), Inches(1.0),
         [(say, 13, TEXT, False), (move, 11, MUTED, False)], anchor=MSO_ANCHOR.MIDDLE, space=2)
    y = y + Inches(1.25)
notes(s, "Slow down on Sentinel — that's the moment of the talk. Use the pause after DEPLOYMENT BLOCKED.")

# ---------------------------------------------------------------- Slide 11: closing + tactics
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "3 · SPEAKER NOTES", "Close & engagement tactics")
box(s, Inches(0.6), Inches(1.9), Inches(12.1), Inches(1.9), fill=CARD)
text(s, Inches(0.9), Inches(2.05), Inches(11.5), Inches(1.6),
     [("\u201cFour teammates: one clarified, one proposed, one caught the flaw, one held the gate — "
       "and a human made the call. Not AI replacing engineers. AI accelerating delivery while humans "
       "own accountability.\u201d", 18, TEXT, True)], anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(0.6), Inches(4.0), Inches(12), Inches(0.5), [("Leave them with: \u201cAI proposes. Humans dispose.\u201d", 20, ACCENT, True)])
bullets(s, Inches(0.6), Inches(4.8), Inches(12), Inches(2.2),
        [("Ask before you show — pose the question, THEN reveal Sentinel catching it", 0, TEXT, False),
         ("Use the pause — after DEPLOYMENT BLOCKED, stop talking for 3 seconds, then vote", 0, TEXT, False),
         ("Name them like people — \u201cRhea noticed…\u201d, not \u201cthe requirements agent\u201d", 0, TEXT, False)], size=16, space=12)
notes(s, "Open the Audit Trail at the end — every step recorded, ending with a human's approval.")

# ---------------------------------------------------------------- Slide 12: reference
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "REFERENCE", "45-min flow & scenario outcomes")
bullets(s, Inches(0.6), Inches(1.9), Inches(6.0), Inches(5),
        [("0–5  Problem & agentic SDLC concept", 0, TEXT, False),
         ("5–12  The four agents + human authorization", 0, TEXT, False),
         ("12–20  Spring AI: ChatModel, ChatClient, structured output", 0, TEXT, False),
         ("20–27  Run the pipeline; events, tools, typed results", 0, TEXT, False),
         ("27–32  Human clarification & governance", 0, TEXT, False),
         ("32–40  Exercise: add the Confluence Agent", 0, ACCENT, True),
         ("40–45  AFTER comparison & recap", 0, TEXT, False)], size=16, space=12)
box(s, Inches(6.9), Inches(1.9), Inches(5.8), Inches(4.9), fill=CARD)
bullets(s, Inches(7.1), Inches(2.05), Inches(5.4), Inches(4.6),
        [("NORMAL → approve → DEPLOYED", 0, TEXT, False),
         ("SECURITY_FAILURE → revise/review → DEPLOYED", 0, TEXT, False),
         ("AMBIGUOUS_REQUIREMENT → pauses", 0, TEXT, False),
         ("TEST_FAILURE → BLOCKED", 0, TEXT, False),
         ("MISSING_APPROVAL → WAITING_FOR_APPROVAL", 0, TEXT, False),
         ("HALLUCINATED_API → BLOCKED", 0, TEXT, False),
         ("PROMPT_INJECTION → ignored → approval", 0, TEXT, False),
         ("Guarantees apply to DEMO only; LIVE uses the real provider", 0, MUTED, False)], size=14, space=10)
notes(s, "Scenario guarantees apply to DEMO only. Tests and deployment are simulated; nothing ships "
         "for real. Never promise a deterministic LIVE security finding.")

out = Path(__file__).resolve().parent / "enterprise-copilot-briefing.pptx"
prs.save(str(out))
print("Saved:", out)
